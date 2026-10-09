from pathlib import Path
import json, re, zipfile, shutil
import pandas as pd
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path('/mnt/data/apcvp2026_pci_deescalation'); FIG=BASE/'figures'
summary=json.loads((BASE/'Analysis_Summary.json').read_text())
sc=pd.read_csv(BASE/'Scenario_Results.csv'); pr=pd.read_csv(BASE/'Global_Sensitivity_PRCC.csv'); mc=pd.read_csv(BASE/'Synthetic_Model_Comparison_AIC.csv'); rob=pd.read_csv(BASE/'Numerical_Robustness.csv'); st=pd.read_csv(BASE/'Stochastic_Perturbation_Results.csv'); surf=pd.read_csv(BASE/'Bleeding_Ischemia_Risk_Surface.csv')

TITLE='Dynamical Modelling of Bleeding–Ischemia Risk Surfaces for Antithrombotic De-escalation after PCI Based on Hemoglobin, Renal Function, and Prior Bleeding'
ALT2='Dynamical Modelling of Competing Bleeding–Ischemia Hazards for Post-PCI Antithrombotic De-escalation Using Hemoglobin, eGFR, and Prior Bleeding'
ALT3='Dynamical Modelling of Antithrombotic Switching Boundaries after PCI from Hemoglobin, Renal Function, Prior Bleeding, and Time-Dependent Ischemic Hazard'
KEYWORDS='Percutaneous Coronary Intervention; Antithrombotic De-escalation; Competing Risk; Dynamical Modelling; Optimal Stopping; Uncertainty Quantification'

sources=[
(1,2023,'ESC Guidelines for the management of acute coronary syndromes','Clinical timing and de-escalation logic','De-escalation may be considered beyond 30 days in selected ACS patients to reduce bleeding; first-30-day de-escalation is not recommended.','https://academic.oup.com/ehjacc/article/13/1/55/7280662'),
(2,2019,'Urban P et al. ARC-HBR consensus','Hemoglobin, eGFR, prior bleeding thresholds','Hb <11 g/dL and eGFR <30 mL/min are major HBR criteria; eGFR 30–59 and mild anemia are minor criteria; prior spontaneous bleeding is a criterion.','https://pmc.ncbi.nlm.nih.gov/articles/PMC6736433/'),
(3,2017,'Costa F et al. PRECISE-DAPT derivation and validation','Variable selection','Age, hemoglobin, creatinine clearance, white-cell count, and previous spontaneous bleeding form the five-item PRECISE-DAPT score.','https://pubmed.ncbi.nlm.nih.gov/28290994/'),
(4,2022,'Gragnano F et al. PRECISE-DAPT in GLOBAL LEADERS and GLASSY','External bleeding-risk context','PRECISE-DAPT remained informative under DAPT and ticagrelor-monotherapy strategies.','https://pubmed.ncbi.nlm.nih.gov/32941620/'),
(5,2016,'Baber U et al. PARIS risk scores','Dual-risk renal signal','Renal dysfunction predicted both coronary thrombotic events and major bleeding; anemia predicted major bleeding.','https://pubmed.ncbi.nlm.nih.gov/27079334/'),
(6,2016,'Yeh RW et al. DAPT prediction rule','Benefit–risk decision precedent','Extended thienopyridine reduced ischemia but increased bleeding; the study formalized a benefit–risk difference after PCI.','https://jamanetwork.com/journals/jama/fullarticle/2508253'),
(7,2017,'Cuisset T et al. TOPIC','De-escalation event-rate anchor','At 1 year, the composite endpoint was 13.4% vs 26.3%; BARC ≥2 bleeding 4.0% vs 14.9% after switching vs unchanged therapy.','https://pubmed.ncbi.nlm.nih.gov/28510646/'),
(8,2017,'Sibbing D et al. TROPICAL-ACS','Guided de-escalation precedent','Platelet-function-guided prasugrel-to-clopidogrel de-escalation was non-inferior for net clinical benefit.','https://pubmed.ncbi.nlm.nih.gov/28855078/'),
(9,2019,'Mehran R et al. TWILIGHT','Bleeding/ischemia anchor','After 3 months of DAPT, BARC 2/3/5 bleeding was 4.0% with ticagrelor monotherapy vs 7.1% with ticagrelor plus aspirin; death/MI/stroke was 3.9% in both groups.','https://www.nejm.org/doi/full/10.1056/NEJMoa1908419'),
(10,2019,'Hahn JY et al. SMART-CHOICE','Bleeding/ischemia anchor','At 12 months, MACCE was 2.9% vs 2.5% and bleeding 2.0% vs 3.4% for P2Y12 monotherapy vs DAPT.','https://jamanetwork.com/journals/jama/fullarticle/2736564'),
(11,2019,'Watanabe H et al. STOPDAPT-2','Short-DAPT anchor','One-month DAPT followed by clopidogrel monotherapy produced a 2.4% vs 3.7% composite cardiovascular/bleeding endpoint at 1 year.','https://jamanetwork.com/journals/jama/fullarticle/2736563'),
(12,2020,'Kim HS et al. HOST-REDUCE-POLYTECH-ACS','Dose de-escalation anchor','Prasugrel dose de-escalation reduced net adverse clinical events (7.2% vs 10.1%) mainly through less bleeding without increased ischemic risk.','https://pubmed.ncbi.nlm.nih.gov/32882163/'),
(13,2021,'Kim CJ et al. TALOS-AMI','Primary literature event-rate anchor','From month 1 to 12, net events were 4.6% vs 8.2%, ischemic composite 2.1% vs 3.1%, and BARC 2/3/5 bleeding 3.0% vs 5.6%.','https://pubmed.ncbi.nlm.nih.gov/34627490/'),
(14,2021,'Valgimigli M et al. MASTER DAPT','HBR abbreviated-DAPT anchor','Major or clinically relevant nonmajor bleeding was 6.5% vs 9.4% with abbreviated vs standard therapy, while major cardiac/cerebral outcomes were non-inferior.','https://www.nejm.org/doi/full/10.1056/NEJMoa2108749'),
(15,2011,'Mehran R et al. BARC consensus','Outcome definition','Standardized Bleeding Academic Research Consortium definitions for cardiovascular trials.','https://pubmed.ncbi.nlm.nih.gov/21670242/'),
(16,2016,'Anemia after PCI systematic review and meta-analysis','Hemoglobin biological rationale','Across 17 studies and 68,528 patients, anemia was associated with increased long-term ischemic and bleeding events.','https://pubmed.ncbi.nlm.nih.gov/26716044/'),
(17,2019,'PARIS anemia analysis','Hemoglobin and treatment-disruption rationale','Baseline anemia was associated with higher adjusted MACE and major bleeding after PCI.','https://pubmed.ncbi.nlm.nih.gov/30998384/'),
]
src=pd.DataFrame(sources,columns=['ID','Year','Source','Used_for','Anchor_or_role','URL']);src.to_csv(BASE/'Literature_Parameter_Sources.csv',index=False)

params=[
('Hb','g/dL','Observed clinical variable','ARC-HBR Hb <11 g/dL major HBR criterion','Literature-derived clinical anchor'),
('eGFR','mL/min/1.73m2','Observed clinical variable','ARC-HBR <30 major; 30–59 minor HBR','Literature-derived clinical anchor'),
('P','0/1','Prior spontaneous bleeding history','ARC-HBR/PRECISE-DAPT','Literature-derived variable definition'),
('t0','days','Earliest de-escalation landmark','30 days','ESC 2023 boundary'),
('b0','annualized hazard scale','Baseline bleeding hazard scale','0.026','Provisional structural coefficient calibrated only to keep simulations within published trial-scale ranges'),
('i0','annualized hazard scale','Baseline ischemic hazard scale','0.090','Provisional structural coefficient'),
('etaB','dimensionless','Potent-therapy bleeding amplification','0.70','Provisional structural coefficient'),
('etaI','dimensionless','Potent-therapy ischemic protection','1.15','Provisional structural coefficient'),
('deltaI','per month','Post-landmark ischemic hazard decay','0.180','Provisional structural coefficient'),
('lambdaB','utility weight','Bleeding loss weight','1.00','Illustrative decision weight'),
('lambdaI','utility weight','Ischemic loss weight','1.70','Illustrative decision weight'),
('cs','utility','Switching penalty','0.003','Illustrative decision cost'),
]
pd.DataFrame(params,columns=['Parameter','Unit','Meaning','Reference_value_or_anchor','Evidence_status']).to_csv(BASE/'Model_Parameters.csv',index=False)

EQS=[
('Hemoglobin vulnerability','q_Hb = [1 + exp{(Hb - 11.5)/0.85}]^(-1)'),
('Renal vulnerability','q_G = [1 + exp{(eGFR - 50)/14}]^(-1)'),
('Prior-bleeding state','P ∈ {0,1}'),
('Treatment intensity','u(t) = 1 for potent DAPT; u(t) = 0.55 for the illustrative de-escalated state'),
('Bleeding hazard','h_B(t) = b_0 exp{β_H q_Hb + β_G q_G + β_P P + η_B u(t)}'),
('Ischemic hazard','h_I(t) = i_0 exp{γ_H q_Hb + γ_G q_G + γ_P P - η_I u(t) - δ_I t}'),
('Latent bleeding burden','dB/dt = k_B h_B(t)[1-B(t)] - r_B B(t)'),
('Latent ischemic burden','dI/dt = k_I h_I(t)[1-I(t)] - r_I I(t)'),
('Net competing-risk margin','Ψ(t) = λ_B[B(t)+h_B(t)] - λ_I[I(t)+h_I(t)]'),
('Latent regime','Z(t) ∈ {Ischemic-dominant, Balanced, Bleeding-dominant, De-escalation-favorable}'),
('Transition probability','Pr[Z(t+Δt)=j | Z(t)=i,x(t)] = [exp{Q(x(t))Δt}]_ij'),
('Cumulative bleeding probability','P_B(T) = 1 - exp{-∫_0^T h_B(s) ds}'),
('Cumulative ischemic probability','P_I(T) = 1 - exp{-∫_0^T h_I(s) ds}'),
('Expected clinical loss','L(τ) = λ_B P_B(T|τ) + λ_I P_I(T|τ) + c_s 1(τ<T)'),
('Optimal stopping time','τ* = arg min_{τ∈[0,T]} L(τ)'),
('Stochastic perturbation','dX_t = f(X_t,θ)dt + Σ dW_t'),
('Synthetic AIC benchmark','AIC = n log(RSS/n) + 2k'),
('Global sensitivity','PRCC(θ_j,τ*) = Corr[rank-residual(θ_j), rank-residual(τ*)]'),
]
(BASE/'Model_Equations_List.txt').write_text('\n\n'.join([a+'\n'+b for a,b in EQS]))

# DOCX helpers
def setup(doc):
    sec=doc.sections[0];sec.top_margin=Inches(.62);sec.bottom_margin=Inches(.62);sec.left_margin=Inches(.72);sec.right_margin=Inches(.72)
    styles=doc.styles;styles['Normal'].font.name='Arial';styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial');styles['Normal'].font.size=Pt(9.6)
    for nm,sz in [('Title',16),('Heading 1',13),('Heading 2',11.5),('Heading 3',10.5)]:
        st=styles[nm];st.font.name='Arial';st._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial');st.font.size=Pt(sz);st.font.bold=True
    return doc

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),fill);tcPr.append(shd)

def table(doc,df,fontsize=7.2):
    t=doc.add_table(rows=1,cols=len(df.columns));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,c in enumerate(df.columns):
        cell=t.rows[0].cells[j];cell.text=str(c);shade(cell,'1F4E78')
        for rr in cell.paragraphs[0].runs: rr.font.bold=True;rr.font.color.rgb=RGBColor(255,255,255);rr.font.size=Pt(fontsize)
    for i,row in df.iterrows():
        cells=t.add_row().cells
        for j,v in enumerate(row):
            cells[j].text=str(v);cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for rr in cells[j].paragraphs[0].runs:rr.font.size=Pt(fontsize)
            if i%2:shade(cells[j],'F3F6F9')
    return t

def para(doc,text,bold=False,italic=False):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=1.05;r=p.add_run(text);r.bold=bold;r.italic=italic;return p

def eqline(doc,label,eq):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(4);r=p.add_run(label+'   ');r.bold=True;r.font.size=Pt(9.5);r=p.add_run(eq);r.font.name='Cambria Math';r.font.size=Pt(10.5)

def add_fig(doc,name,caption,width=6.65):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(FIG/name),width=Inches(width));cp=doc.add_paragraph();cp.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=cp.add_run(caption);r.italic=True;r.font.size=Pt(8)

def hyperlink(paragraph,text,url):
    part=paragraph.part;rid=part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rid);r=OxmlElement('w:r');rPr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'0563C1');rPr.append(color);u=OxmlElement('w:u');u.set(qn('w:val'),'single');rPr.append(u);r.append(rPr);tx=OxmlElement('w:t');tx.text=text;r.append(tx);h.append(r);paragraph._p.append(h)

# manuscript
D=setup(Document())
p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(TITLE);r.bold=True;r.font.size=Pt(16)
p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('Literature-parameterized computational proof-of-concept for cardiovascular pharmacotherapy research');r.italic=True;r.font.size=Pt(9.5)
D.add_heading('Research Integrity Statement',level=1)
para(D,'This article is a theoretical and computational proof-of-concept. No patient-level registry records, prospective observations, laboratory measurements, or unpublished clinical datasets were analyzed. Published PCI, acute coronary syndrome, bleeding-risk, and antiplatelet-therapy studies were used to define observable variables, guideline landmarks, event-rate ranges, and biologically plausible directions. Hazard weights, latent-state coefficients, utility weights, switching costs, and stochastic terms are provisional structural quantities unless explicitly labeled literature-derived. All numerical findings reported below are reproducible simulation outputs, not validated clinical effect estimates, treatment thresholds, sensitivity/specificity values, or recommendations. Final effect estimates must be recalculated from underlying registry or longitudinal cohort data before scientific or clinical interpretation.')

D.add_heading('Three Title Choices',level=1)
for i,t in enumerate([TITLE,ALT2,ALT3],1): para(D,f'{i}. {t} [{len(re.findall(r"\\b[\\w–-]+\\b",t))} words]')
para(D,'Six iconic keywords: '+KEYWORDS+'.')
para(D,'Proposed conference fit: antithrombotic/cardiovascular pharmacotherapy, PCI/acute coronary syndromes, and bleeding-risk optimization. The APCVP 2026 website currently blocks automated retrieval of the category dropdown; this wording should therefore be treated as a best-fit scientific mapping rather than an asserted official portal label.')

D.add_heading('Structured Abstract',level=1)
D.add_heading('Background/Aim',level=2)
para(D,'Antithrombotic de-escalation after PCI is a recurrent clinical trade-off: sustained platelet inhibition reduces thrombotic risk but exposes patients to bleeding, while the relative importance of both hazards changes with time and patient vulnerability. Current scores identify high-risk patients but do not explicitly represent a moving switching boundary in which the same treatment can carry different net value at different post-PCI states. We developed a literature-parameterized dynamical proof-of-concept using hemoglobin, renal function, prior bleeding, and time since a 30-day post-ACS landmark to construct competing bleeding–ischemia risk surfaces and a theoretical optimal-stopping rule for later clinical calibration.')
D.add_heading('Methods',level=2)
para(D,'Hemoglobin and eGFR were transformed to continuous vulnerability functions centered on clinically recognized ARC-HBR regions, while prior bleeding was binary. A time-varying competing-hazard system linked antithrombotic intensity to bleeding amplification and ischemic protection; latent bleeding and ischemic burdens were propagated with coupled ordinary differential equations and summarized by a state-dependent net-risk margin. Four regimes—ischemic-dominant, balanced, bleeding-dominant, and de-escalation-favorable—were defined for exploratory state-transition analysis. Sequential loss minimization identified the theoretical switching time after the 30-day landmark. Computation used adaptive RK45 with DOP853 verification, a 21×17×2 hemoglobin–eGFR–bleeding-history grid, five clinical stress scenarios, 320-set Latin-hypercube sampling with partial rank correlation coefficients, local elasticity, 600 stochastic vulnerability perturbations, and a four-model synthetic AIC recovery benchmark. No clinical outcome model was fitted.')
D.add_heading('Results',level=2)
ref=summary['reference_profile']; sto=summary['stochastic']; top=summary['top_prcc']
para(D,f"In the reference state (Hb 11.5 g/dL, eGFR 45 mL/min/1.73 m², no prior bleeding), the loss-minimizing switch occurred 3.0 months after the 30-day landmark, corresponding to approximately day 121 after the index event; the model-implied remaining-horizon bleeding and ischemic probabilities at that policy were {100*ref['optimal_switch'][2]:.1f}% and {100*ref['optimal_switch'][3]:.1f}%, respectively. Scenario switching times ranged from immediate at the 30-day landmark for combined HBR vulnerability to 0.6 months for prior bleeding, 2.4 months for isolated anemia, 4.2 months for renal vulnerability, and 4.6 months for the low-vulnerability profile. Under 600 stochastic perturbations, the median theoretical switch was {sto['median_switch_month']:.1f} months after day 30 (2.5th–97.5th percentile {sto['p2_5']:.1f}–{sto['p97_5']:.1f}); {100*sto['immediate_fraction']:.1f}% of simulations selected the 30-day boundary. Global sensitivity ranked treatment-related bleeding amplification (PRCC {top[0]['PRCC_optimal_switch_month']:.3f}), baseline ischemic hazard ({top[1]['PRCC_optimal_switch_month']:.3f}), baseline bleeding hazard ({top[2]['PRCC_optimal_switch_month']:.3f}), and bleeding/ischemic utility weights as the dominant structural drivers. Solver outputs were numerically stable across RK45 and DOP853 settings. In the explicitly synthetic model-recovery experiment, an exponential ischemic-decay model had the lowest AIC; the full structural model was within ΔAIC 2.0, whereas linear and constant hazard alternatives were substantially weaker. These are internal structural results, not clinical discrimination statistics.")
D.add_heading('Conclusions',level=2)
para(D,'The current framework establishes a mathematically explicit structure for a familiar pharmacotherapy problem: when changing bleeding vulnerability, renal function, prior bleeding, and declining post-PCI ischemic pressure alter the balance between continued potent therapy and de-escalation. The proof-of-concept produces a falsifiable switching surface rather than a fixed clinical cutoff and illustrates how financial-engineering tools—competing downside risk, state transitions, first-passage timing, uncertainty propagation, and optimal stopping—can be translated into clinician-facing cardiovascular decision science. Its immediate use is to define variables, sampling times, candidate states, and prespecified validation hypotheses for a retrospective PCI registry analysis; its potential benefit is a future trajectory-aware decision aid that makes the bleeding–ischemia trade-off more transparent. Parameter refinement and clinical calibration are ongoing. Clinical calibration and external validation are required before any use in patient care. This mechanistic proof-of-concept provides a mathematically testable framework for identifying clinically relevant intervention windows that conventional association-based analyses may not reveal. It is designed to guide variable selection, protocol development, and subsequent retrospective or prospective validation in collaboration with interventional cardiology and cardiovascular pharmacotherapy teams.')

D.add_heading('1. Background and Aim',level=1)
para(D,'Dual antiplatelet therapy after PCI embodies a genuine competing-risk problem. The 2023 ESC ACS guideline preserves 12-month potent DAPT as the default in many ACS patients but permits individualized shortening or P2Y12 de-escalation when bleeding risk becomes clinically important, while specifically advising against de-escalation during the first 30 days after ACS. ARC-HBR and PRECISE-DAPT place hemoglobin, renal impairment, and previous bleeding among established bleeding-risk determinants, and PARIS shows that renal dysfunction can carry both thrombotic and bleeding information. Trials including TOPIC, TROPICAL-ACS, TWILIGHT, SMART-CHOICE, HOST-REDUCE-POLYTECH-ACS, TALOS-AMI, and MASTER DAPT demonstrate that reducing antiplatelet intensity can lower bleeding in selected stabilized patients without a uniform ischemic penalty, but these trials do not provide a universal individual switching surface. The present study therefore asks a narrower theoretical question: can a parsimonious model using routinely available hemoglobin, eGFR, prior bleeding, and post-PCI time generate an explicit dynamic boundary between continued potent therapy and candidate de-escalation? The financial-engineering contribution is methodological rather than metaphorical: the problem is framed as sequential risk management under competing losses, changing hazard regimes, uncertainty, and a stopping decision.')

D.add_heading('2. Methods',level=1)
D.add_heading('2.1 Study Design and Evidence Status',level=2)
para(D,'We designed a literature-parameterized theoretical study integrating mathematical biology, computational medicine, cardiovascular pharmacotherapy, quantitative decision science, and financial-engineering methods. Clinical literature supplied the definition and direction of observable variables, the 30-day landmark, bleeding-risk thresholds, and event-rate ranges used as plausibility checks. The literature did not uniquely identify the hazard weights, ODE coefficients, utility weights, switching cost, or stochastic covariance matrix; those quantities are therefore explicitly provisional. The model is a structural hypothesis generator intended for later estimation using individual longitudinal PCI data with treatment exposure, serial hemoglobin/renal measurements, bleeding events, and ischemic events.')
D.add_heading('2.2 Observable Variables and Literature Anchors',level=2)
para(D,'The primary observables were hemoglobin in g/dL, eGFR in mL/min/1.73 m², and prior spontaneous bleeding. ARC-HBR supplied clinically interpretable anchors: hemoglobin <11 g/dL and eGFR <30 mL/min are major high-bleeding-risk criteria, moderate CKD with eGFR 30–59 mL/min is a minor criterion, and previous spontaneous bleeding contributes according to timing and severity. PRECISE-DAPT independently supports hemoglobin, renal function through creatinine clearance, and prior bleeding as core variables. Time was measured from a 30-day post-index landmark so that the proof-of-concept respects the ESC statement that de-escalation within the first 30 days after ACS is not recommended. Published trial event rates were used only to check whether illustrative simulations remained on a clinically recognizable scale, not to create pooled patient-level effect estimates.')
D.add_heading('2.3 Model Formulation',level=2)
for a,b in EQS: eqline(D,a,b)
para(D,'The hemoglobin and renal vulnerability transforms are continuous rather than hard binary cutoffs so that the surface can be differentiated and later estimated from patient-level data. Potent DAPT was assigned u=1 and the illustrative de-escalated state u=0.55 solely as a normalized exposure contrast; this number is not a dose conversion between ticagrelor, prasugrel, and clopidogrel. The ischemic hazard was allowed to decay after the 30-day landmark, reflecting the clinical premise that thrombotic pressure is front-loaded after ACS/PCI, whereas the bleeding hazard remains strongly exposure- and vulnerability-dependent. The latent B(t) and I(t) states provide path dependence and permit a future state-space implementation when repeated clinical measurements become available.')
D.add_heading('2.4 Theoretical Analysis',level=2)
para(D,'Existence and uniqueness follow locally from continuous, locally Lipschitz right-hand sides for finite coefficients. Positivity and boundedness were checked for the latent burden states on [0,1], because inflow terms vanish at the upper boundary and recovery terms vanish at the lower boundary in the implemented parameter domain. Frozen-covariate analysis treats hemoglobin, eGFR, prior bleeding, and treatment intensity as fixed over a short interval and examines the B–I subsystem by Jacobian eigenvalues. A phase-plane representation is therefore meaningful for the latent two-dimensional subsystem, whereas full-system equilibria would be clinically misleading because post-PCI time intentionally changes the ischemic hazard. The principal theoretical object is not a disease equilibrium but the zero level set of the expected-loss difference, which defines a moving switching boundary in hemoglobin–eGFR space conditional on bleeding history and time.')
D.add_heading('2.5 Computational Analysis',level=2)
para(D,'Reference latent trajectories were integrated with adaptive Runge–Kutta RK45 at relative tolerance 10^-9 and absolute tolerance 10^-11 and checked against DOP853 over maximum step sizes from 1.0 to 0.1 month. The switching surface evaluated hemoglobin from 9.5 to 14.5 g/dL in 0.25-g/dL increments and eGFR from 20 to 100 mL/min/1.73 m² in 5-unit increments, separately for patients with and without prior bleeding. At each grid point, 56 candidate switch times over the 11-month post-landmark horizon were evaluated by closed-form integration of the piecewise hazards, and the loss-minimizing time was retained. Five stress scenarios represented low vulnerability, renal vulnerability, anemia, prior bleeding, and combined HBR.')
D.add_heading('2.6 Robustness, Uncertainty, and Sensitivity',level=2)
para(D,'Structural uncertainty was explored with 320 Latin-hypercube draws over prespecified ranges for hazard scales, clinical-vulnerability weights, treatment effects, ischemic decay, and utility weights. Partial rank correlation coefficients linked each parameter to the optimal switching time after residualizing ranks against all remaining sampled parameters. Local elasticities used a +1% perturbation around the reference values. Separately, 600 stochastic clinical profiles perturbed hemoglobin and eGFR around the reference state and sampled prior bleeding, thereby separating uncertainty in patient state from uncertainty in structural coefficients. These stochastic results are scenario dispersion rather than empirically estimated measurement error.')
D.add_heading('2.7 Model Comparison and AIC',level=2)
para(D,'Because no clinical likelihood was fitted, AIC was restricted to an explicitly synthetic model-recovery benchmark. Monthly pseudo-observations of bleeding and ischemic hazard were generated from the reference structure with Gaussian noise and compared across the full competing-hazard model, independent linear hazards, constant hazards, and a reduced exponential ischemic-decay model. AIC was calculated as n log(RSS/n)+2k. The comparison answers only whether the generating temporal structure can be recovered under controlled noise; it is not evidence that the framework outperforms PRECISE-DAPT, ARC-HBR, DAPT score, or any validated clinical prediction model.')
D.add_heading('2.8 Planned Clinical Calibration and Statistical Validation',level=2)
para(D,'The next-stage analysis should use patient-level PCI data and prespecify the index population, ACS/chronic coronary syndrome status, antiplatelet regimen, treatment changes, serial hemoglobin and creatinine/eGFR, prior bleeding, BARC bleeding outcomes, myocardial infarction, stent thrombosis, stroke, and death. Structural parameters could be estimated with joint longitudinal–time-to-event or state-space methods, with competing-risk or multi-state survival models as statistical comparators. Candidate formulations should be compared by AIC/BIC and out-of-sample log likelihood; calibration intercept/slope, time-dependent Brier score, cause-specific and subdistribution hazards, bootstrap optimism correction, and decision-curve net benefit should be reported only after actual data fitting. Missing laboratory measurements, informative follow-up, treatment-confounding, indication bias, and time-dependent exposure require explicit handling. Final effect estimates must be recalculated from the underlying registry data before scientific or clinical interpretation.')

D.add_heading('3. Results',level=1)
D.add_heading('3.1 Literature-Anchored Clinical Plausibility',level=2)
para(D,'The chosen observables and boundary design were concordant with established cardiovascular evidence. ARC-HBR places Hb <11 g/dL and eGFR <30 mL/min among major HBR criteria, while PRECISE-DAPT includes hemoglobin, creatinine clearance, and previous bleeding. The simulated direction of treatment effect was also consistent with the published de-escalation literature: TALOS-AMI reported BARC 2/3/5 bleeding of 3.0% vs 5.6% after de-escalation vs continued ticagrelor-based DAPT; TWILIGHT reported 4.0% vs 7.1% BARC 2/3/5 bleeding after aspirin withdrawal; MASTER DAPT reported 6.5% vs 9.4% major or clinically relevant nonmajor bleeding with abbreviated vs standard therapy. These values were not pooled or fitted as individual outcomes; they were used as plausibility envelopes for simulated risks.')
D.add_heading('3.2 Dynamic Switching Surface',level=2)
para(D,'The theoretical switching boundary varied smoothly with hemoglobin and eGFR and shifted earlier when previous bleeding was present. For the reference state of Hb 11.5 g/dL and eGFR 45 mL/min/1.73 m² without previous bleeding, the minimum expected loss occurred 3.0 months after the 30-day landmark, approximately day 121 post-index. In the low-vulnerability profile (Hb 13.8, eGFR 85, no prior bleeding), the optimum shifted to 4.6 months after the landmark, whereas isolated anemia (Hb 10.8) shifted it to 2.4 months. Previous bleeding at otherwise preserved laboratory values moved the minimum to 0.6 months after the landmark, and combined Hb 10.5 g/dL, eGFR 25, and previous bleeding selected the earliest permitted boundary at day 30. These are model-implied first-passage decisions, not treatment recommendations.')
add_fig(D,'Figure_2_Switching_Surface_PB0.png','Figure 1. Hemoglobin–eGFR switching surface without prior bleeding. The color field denotes the simulation-derived loss-minimizing month after the 30-day landmark, not a clinical cutoff.')
add_fig(D,'Figure_3_Switching_Surface_PB1.png','Figure 2. The same theoretical surface conditional on prior bleeding, demonstrating systematic leftward movement of the switching boundary.')
D.add_heading('3.3 Scenario Stress Tests',level=2)
ss=sc[['profile','Hb_g_dL','eGFR','prior_bleeding','optimal_switch_month_after_30d','bleed_prob_potent','ischemic_prob_potent','bleed_prob_deesc','ischemic_prob_deesc']].copy()
for c in ['bleed_prob_potent','ischemic_prob_potent','bleed_prob_deesc','ischemic_prob_deesc']:ss[c]=(100*ss[c]).round(2)
ss.columns=['Profile','Hb','eGFR','Prior bleeding','Switch month','Bleed % potent','Ischemia % potent','Bleed % de-escalated','Ischemia % de-escalated']
table(D,ss,6.6)
para(D,'Across the five prespecified scenarios, fully de-escalated exposure over the remaining horizon lowered model-implied bleeding probability by 1.38–7.30 percentage points but increased ischemic probability by 0.81–1.67 percentage points, illustrating why a single universal rule is mathematically inappropriate. The sequential optimizer did not simply choose the lower-bleeding strategy for everyone; it delayed switching in lower-vulnerability profiles because early ischemic loss was weighted more heavily and decayed over time. In the combined HBR scenario, the model-implied potent-DAPT bleeding probability was 31.10% versus 23.80% under the de-escalated exposure, a deliberately high stress-test range reflecting stacked vulnerability rather than a population estimate.')
add_fig(D,'Figure_1_Sequential_Loss_Curves.png','Figure 3. Sequential expected-loss curves under five vulnerability profiles. Minima move earlier as bleeding vulnerability accumulates.')
add_fig(D,'Figure_7_Scenario_Bleeding.png','Figure 4. Scenario stress test of model-implied bleeding probability under potent and de-escalated exposure states.')
D.add_heading('3.4 Latent-State Dynamics and Regime Geometry',level=2)
para(D,'In the reference latent simulation, ischemic burden declined as the explicitly time-dependent ischemic hazard decayed, while bleeding burden remained sustained by hemoglobin/renal vulnerability and continuing treatment exposure. The net margin therefore moved from ischemic-dominant toward balanced and then bleeding-dominant territory, demonstrating the intended mechanism by which a treatment decision can change even if baseline covariates are fixed. This path dependence is the central distinction from a static score: the same Hb and eGFR pair can map to a different decision region at month 2 than at month 6 because the competing hazard geometry has changed.')
add_fig(D,'Figure_4_Reference_Dynamics.png','Figure 5. Reference competing-risk dynamics showing latent bleeding burden, ischemic burden, and the time-varying net margin.')
D.add_heading('3.5 Global Sensitivity and Uncertainty Quantification',level=2)
para(D,f"Across 320 Latin-hypercube structural draws, the strongest determinant of switching time was treatment-associated bleeding amplification (PRCC {top[0]['PRCC_optimal_switch_month']:.3f}), followed by baseline ischemic hazard ({top[1]['PRCC_optimal_switch_month']:.3f}), baseline bleeding hazard ({top[2]['PRCC_optimal_switch_month']:.3f}), the bleeding loss weight ({top[3]['PRCC_optimal_switch_month']:.3f}), and the ischemic loss weight ({top[4]['PRCC_optimal_switch_month']:.3f}). Renal influence on ischemic and bleeding hazards also remained material (|PRCC| approximately {abs(top[5]['PRCC_optimal_switch_month']):.3f} and {abs(top[6]['PRCC_optimal_switch_month']):.3f}). Under 600 stochastic clinical-profile perturbations, median switching occurred 2.8 months after the 30-day landmark with a 2.5th–97.5th percentile range of 0.0–4.0 months; 14.7% selected the earliest permitted switch. The spread demonstrates that uncertainty belongs in the decision surface rather than being hidden behind a single deterministic cutoff.")
add_fig(D,'Figure_5_Global_Sensitivity_PRCC.png','Figure 6. PRCC sensitivity of the theoretical optimal switching month to structural parameters.')
add_fig(D,'Figure_6_Stochastic_Switching_Time.png','Figure 7. Distribution of theoretical switching times under stochastic perturbation of hemoglobin, eGFR, and prior-bleeding status.')
D.add_heading('3.6 Numerical Robustness and Internal Model Recovery',level=2)
spreadB=rob.terminal_B.max()-rob.terminal_B.min();spreadI=rob.terminal_I.max()-rob.terminal_I.min();spreadM=rob.terminal_margin.max()-rob.terminal_margin.min()
para(D,f'Solver verification showed negligible numerical dispersion across RK45 and DOP853 settings: the terminal latent bleeding burden varied by {spreadB:.2e}, ischemic burden by {spreadI:.2e}, and net margin by {spreadM:.2e}. In the synthetic recovery benchmark, the reduced exponential ischemic-decay model had the lowest AIC; the full structural model remained within ΔAIC 2.00, while the independent linear and constant hazard models had ΔAIC {mc.loc[mc.model=="Independent linear hazard model","delta_AIC"].iloc[0]:.2f} and {mc.loc[mc.model=="Constant competing-hazard model","delta_AIC"].iloc[0]:.2f}. This result argues for temporal nonlinearity but does not establish clinical superiority of the full model; the correct model will have to be selected using real longitudinal likelihoods.')
add_fig(D,'Figure_8_Synthetic_AIC.png','Figure 8. Synthetic AIC comparison used solely for internal structural recovery, not clinical validation.')
D.add_heading('3.7 Clinical Interpretation',level=2)
para(D,'The principal insight is not that hemoglobin, renal function, or previous bleeding are new risk factors; they are already established. The testable hypothesis is that these variables define a time-dependent geometry in which the post-PCI treatment decision is path dependent. Renal dysfunction is particularly informative because it can increase both ischemic and bleeding risk, preventing simplistic “higher bleeding risk means immediate de-escalation” logic. A conventional regression can estimate interactions, but the dynamical framework additionally encodes hazard decay, latent burden, sequential updating, uncertainty propagation, and the first-passage time to a decision boundary. A future registry study can directly test whether this additional structure improves calibration, net benefit, and treatment-policy transportability beyond established scores.')

D.add_heading('4. Discussion',level=1)
para(D,'This proof-of-concept translates a familiar interventional-cardiology dilemma into a formal competing-risk decision system. Its novelty is not the discovery of Hb, eGFR, or prior bleeding as predictors, and it does not claim a new bedside score. The contribution is a falsifiable mathematical object—the switching surface—that changes with patient vulnerability and elapsed post-PCI time. This permits questions that a static association model does not naturally ask: how far is the current state from a boundary, how rapidly is that distance changing, how uncertain is the boundary, and how does a treatment decision alter future bleeding and ischemic trajectories? The framework also provides an explicit bridge between financial engineering and cardiovascular pharmacotherapy: downside-risk weighting becomes clinical loss, regime switching becomes a latent treatment-risk state, stochastic perturbation becomes uncertainty propagation, and optimal stopping becomes the choice of when a candidate de-escalation policy first minimizes expected harm. These translations are mathematical tools, not claims that patients behave like financial assets.')
D.add_heading('4.1 Strengths',level=2)
para(D,'The framework uses three clinically accessible variables with strong literature support; respects the guideline-defined early hazard period by starting the decision problem after day 30; separates literature-derived anchors from provisional coefficients; exposes rather than hides sensitivity to utility weights; and produces fully reproducible code, grids, scenario outputs, uncertainty simulations, and figures. It is computationally light enough to reproduce within minutes and can therefore be rapidly re-estimated when registry data become available.')
D.add_heading('4.2 Limitations',level=2)
para(D,'The numerical outputs remain simulation-derived. Treatment intensity is normalized rather than drug-specific, serial hemoglobin and eGFR changes are not yet jointly modeled, prior bleeding is simplified, and competing death, recurrent revascularization, oral anticoagulation, procedural complexity, diabetes, age, sex, platelet count, and genotype are omitted from the core proof-of-concept. Published trials differ in populations, drugs, timing, and endpoints; their rates therefore cannot be treated as exchangeable calibration data. The utility weights and switching cost are intentionally illustrative. Real clinical use would require causal handling of treatment switching and confounding by indication, external validation, and prospective impact assessment.')

D.add_heading('5. Conclusions',level=1)
para(D,'Dynamical modelling can convert post-PCI bleeding–ischemia balance from a static risk label into a time-dependent, testable switching surface. In illustrative simulations, lower hemoglobin, impaired renal function, and previous bleeding shifted the loss-minimizing boundary earlier, while lower vulnerability supported delayed switching as early ischemic pressure decayed. The usage of the present framework is methodological and translational: it can guide registry variable selection, measurement schedules, candidate treatment states, and prespecified hypotheses for retrospective calibration. Its potential clinical benefit is a future transparent decision-support layer that displays how bleeding and ischemic liabilities evolve together rather than reporting isolated risks. Its research advantage is reproducibility, explicit uncertainty, mathematical falsifiability, and direct comparability with established risk scores and standard survival models. Parameter refinement and biological calibration are ongoing. Clinical calibration and external validation are planned using retrospective or prospective PCI cohorts. This mechanistic proof-of-concept provides a mathematically testable framework for identifying clinically relevant intervention windows that conventional association-based analyses may not reveal. It is designed to guide variable selection, protocol development, and subsequent retrospective or prospective validation in collaboration with interventional cardiovascular teams. Final effect estimates must be recalculated from the underlying registry data before scientific or clinical interpretation.')

D.add_heading('Data, Code, and Reproducibility',level=1)
para(D,'All simulation-derived numerical results in this manuscript can be regenerated from the accompanying Python code. The package includes the hemoglobin–eGFR switching surface, scenario results, Latin-hypercube parameter samples, sensitivity outputs, stochastic perturbations, numerical solver checks, synthetic AIC benchmark, parameter provenance, figures, and model-equation list. No patient-level data are redistributed.')

D.add_heading('References and Parameter Sources',level=1)
for _,r in src.iterrows():
    p=D.add_paragraph();p.paragraph_format.space_after=Pt(3);p.add_run(f"{int(r.ID)}. {r.Source} ({int(r.Year)}). ").bold=True;p.add_run(r.Anchor_or_role+' ');hyperlink(p,'Source / access page',r.URL)

D.add_heading('Supplementary Table S1. Model Parameters and Evidence Status',level=1)
table(D,pd.read_csv(BASE/'Model_Parameters.csv'),6.4)
D.add_heading('Supplementary Table S2. Literature and CSV Provenance',level=1)
table(D,src[['ID','Year','Used_for','Anchor_or_role']],6.2)
D.add_heading('Supplementary Table S3. Global Sensitivity',level=1)
table(D,pr.head(15).round(4),6.6)
D.add_heading('Supplementary Table S4. Synthetic AIC',level=1)
table(D,mc.round(4),6.6)

man=BASE/'PCI_Dynamical_Deescalation_Full_Manuscript.docx';D.save(man)

# equations docx
E=setup(Document());p=E.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('Model Equations');r.bold=True;r.font.size=Pt(16);para(E,TITLE,italic=True);para(E,'Equations are intentionally listed one per row. Clinical anchors are literature-informed; structural coefficients remain provisional until patient-level calibration.')
for a,b in EQS:eqline(E,a,b)
E.save(BASE/'Model_Equations_List.docx')

# provenance docx
P=setup(Document());p=P.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('Parameter Sources and CSV Provenance');r.bold=True;r.font.size=Pt(16);para(P,TITLE,italic=True);para(P,'Published studies supply clinical variable definitions, guideline landmarks, observed event-rate ranges, and biological directions. They do not identify every model coefficient. Provisional quantities must be estimated from observed longitudinal PCI data before clinical interpretation.')
P.add_heading('Clinical and Structural Parameters',level=1);table(P,pd.read_csv(BASE/'Model_Parameters.csv'),6.4)
P.add_heading('Literature Sources',level=1)
for _,r in src.iterrows():
    p=P.add_paragraph();p.add_run(f"{int(r.ID)}. {r.Source} ({int(r.Year)}). ").bold=True;p.add_run('Use: '+r.Used_for+'. '+r.Anchor_or_role+' ');hyperlink(p,'Open source/access page',r.URL)
P.save(BASE/'Parameter_Sources_and_CSV_Provenance.docx')

# figure portfolio
F=setup(Document());p=F.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('APCVP 2026 PCI Figure Portfolio');r.bold=True;r.font.size=Pt(16);para(F,TITLE,italic=True)
for fn,cap in [
('Figure_1_Sequential_Loss_Curves.png','Sequential expected-loss curves'),('Figure_2_Switching_Surface_PB0.png','Switching surface without prior bleeding'),('Figure_3_Switching_Surface_PB1.png','Switching surface with prior bleeding'),('Figure_4_Reference_Dynamics.png','Reference competing-risk dynamics'),('Figure_5_Global_Sensitivity_PRCC.png','Global sensitivity'),('Figure_6_Stochastic_Switching_Time.png','Stochastic switching-time distribution'),('Figure_7_Scenario_Bleeding.png','Scenario bleeding stress test'),('Figure_8_Synthetic_AIC.png','Synthetic AIC recovery benchmark')]:
    add_fig(F,fn,cap,6.5)
F.save(BASE/'Figure_Portfolio.docx')

# ZIP package
zip_path=BASE/'APCVP2026_PCI_Dynamical_Modelling_Package.zip'
with zipfile.ZipFile(zip_path,'w',zipfile.ZIP_DEFLATED) as z:
    for p in BASE.rglob('*'):
        if p.is_file() and p!=zip_path:
            z.write(p,p.relative_to(BASE))
print(man)
print(zip_path)
